#!/usr/bin/env python3
"""INIC 4 - Disponibilidad operativa (% de respaldo) de un mes y su Excel de soporte.

% Respaldo = (T_red - T_solar) / T_red x 100, consolidado como promedio ponderado por tiempo
(suma de T_respaldo / suma de T_red) de los eventos de red con T_red mayor a 2 minutos.

Uso:
  python3 om-agency/scripts/indicador_inic4.py --out /tmp/INIC4.xlsx            # mes anterior
  python3 om-agency/scripts/indicador_inic4.py --mes 2026-09 --out /tmp/INIC4.xlsx
Variables de entorno: METRUM_API_URL, METRUM_USERNAME, METRUM_PASSWORD (nunca se escriben en archivos).
Ademas separa el tiempo sin respaldo segun el SOC al apagarse la casa (regla del 05/10/2026): bateria sin carga
(consumo del cliente) vs falla o demora de transicion (reporte_diario.clasificar_evento); el INIC 4 sigue siendo el
indicador definido, y el "INIC 4 sin perdidas por bateria sin carga" se muestra solo como referencia.
Alcance: todos los sistemas de Metrum salvo Piloto Promigas, Piloto Huawei, Castellana Real y los de
om-agency/data/inic4_config.json (excluir_sistemas). Los sistemas que ingresaron a operar en el mes o que
estan solo con baterias cuentan con los eventos posteriores a su primer dato y se rotulan en Observacion.
"""
import argparse, datetime as dt, json, os, subprocess, sys, shutil, tempfile, time
from concurrent.futures import ThreadPoolExecutor
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reporte_diario as R
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

HERE = os.path.dirname(os.path.abspath(__file__))
BOG = R.BOG
ORDEN = ["Prado Verde Primavera", "Reservas de Pance", "Terra by Kaia", "Llanos de Pance", "Porton de la Rivera", "Los Abedules", "Terranova", "Gerona Club House"]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def zn(z):
    return z.title().replace("Condominio ", "").replace(" De ", " de ").replace(" La ", " la ").replace(" By ", " by ")


def primer_dato(did, key):
    r = R.get(f"/api/plugins/telemetry/DEVICE/{did}/values/timeseries?keys={key}&startTs=0&endTs={int(time.time()*1000)}&limit=1&orderBy=ASC&agg=NONE")
    v = r.get(key, [])
    return v[0]["ts"] if v else None


def calcular(h, a, b):
    hoy0 = dt.datetime.now(BOG).replace(hour=0, minute=0, second=0, microsecond=0)
    r = R.procesar(h, a, b, hoy0, [hoy0])
    pd = []
    for tipo, key in (("red", "voltageA"), ("solar", "voltageA"), (None, "BattSOC")):
        ts = [primer_dato(x["id"], key) for x in h["hijos"] if x["mettype"] == tipo]
        ts = [t for t in ts if t]
        if ts:
            pd.append(min(ts))
    r["primer_dato"] = max(pd) if pd else None
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--mes", help="AAAA-MM; por defecto el mes anterior")
    ap.add_argument("--out", default="INIC4.xlsx")
    args = ap.parse_args()
    hoy = dt.datetime.now(BOG)
    if args.mes:
        y, m = map(int, args.mes.split("-"))
    else:
        y, m = (hoy.year, hoy.month - 1) if hoy.month > 1 else (hoy.year - 1, 12)
    ini = dt.datetime(y, m, 1, tzinfo=BOG)
    fin = dt.datetime(y + (m == 12), 1 if m == 12 else m + 1, 1, tzinfo=BOG)
    a, b = ini.timestamp() * 1000, fin.timestamp() * 1000
    cfg = json.load(open(os.path.join(HERE, "..", "data", "inic4_config.json"), encoding="utf-8"))
    excl, umbral = cfg.get("excluir_sistemas", {}), cfg.get("umbral_evento_s", 120)
    R.login()
    flota = R.cargar_flota()
    print(f"Flota: {len(flota)} sistemas; mes {y}-{m:02d}", file=sys.stderr)
    with ThreadPoolExecutor(6) as ex:
        res = list(ex.map(lambda h: calcular(h, a, b), flota))
    inc = [r for r in res if r["casa"] not in excl]
    pref = {z: i for i, z in enumerate(ORDEN)}
    inc.sort(key=lambda r: (pref.get(zn(r["zona"]), 99), zn(r["zona"]), len(r["casa"]), r["casa"]))
    ev = [(r, e) for r in inc for e in r["eventos"] if e["or_s"] > umbral and e["cl_s"] is not None]
    ev.sort(key=lambda x: x[1]["ini"])
    cau = {id(e): R.clasificar_evento(e) for r, e in ev}      # segundos sin respaldo por causa de cada evento
    NOMBRE_CAUSA = {"soc_bajo": "Batería sin carga", "inicio": "Transición a OFF-GRID", "medio": "Falla durante el corte", "fin": "Retorno a la red", "previo": "No atribuible", "sin_dato": "No atribuible"}

    def obs(r):
        sin = r["casa"] in R.SIN_PANELES or (r.get("gen") is not None and r["gen"] < 0)
        nuevo = r["primer_dato"] and r["primer_dato"] > a + 24 * 3600e3
        d = dt.datetime.fromtimestamp(r["primer_dato"] / 1000, BOG).strftime("%d/%m") if r["primer_dato"] else ""
        if sin and nuevo: return f"Solo con batería; datos desde el {d}"
        if sin: return "Solo con batería (sin paneles)"
        if nuevo: return f"Ingresó a operar en el mes; datos desde el {d}"
        return ""

    # ---- Excel
    A = "Arial"
    hf = Font(name=A, bold=True, color="FFFFFF", size=10); hfill = PatternFill("solid", fgColor="1F3A5F")
    nf = Font(name=A, size=10); bf = Font(name=A, size=10, bold=True); inp = Font(name=A, size=10, color="0000FF")
    thin = Side(style="thin", color="C9D1CE"); bd = Border(top=thin, bottom=thin, left=thin, right=thin)
    f = lambda t: dt.datetime.fromtimestamp(t / 1000, BOG).replace(tzinfo=None)
    wb = Workbook(); wr = wb.active; wr.title = "Resumen"; we = wb.create_sheet("Eventos"); ws = wb.create_sheet("Por sistema")
    nE, nS = len(ev) + 1, len(inc) + 1
    we.append(["#", "Sistema", "Zona", "Marca", "Inicio del corte de red", "Fin del corte de red", "Cortes agrupados", "T_red (min)", "T_solar (min)", "T_respaldo (min)", "% respaldo del evento", "SOC al inicio (%)", "SOC mínimo en el corte (%)",
               "Sin respaldo: batería sin carga (min)", "Sin respaldo: demora o falla de transición a OFF-GRID (min)", "Sin respaldo: falla durante el corte (min)", "Sin respaldo: demora en el retorno a la red (min)", "Sin respaldo: no atribuible (min)", "Causa principal"])
    for i, (r, e) in enumerate(ev, start=2):
        we.append([i - 1, r["casa"], zn(r["zona"]), r["marca"], f(e["ini"]), f(e["fin"]), e["n"], e["or_s"] / 60, e["cl_s"] / 60, f"=MAX(0,H{i}-I{i})", f"=J{i}/H{i}", e["soc"], e["soc_dur"]] + [cau[id(e)][k] / 60 for k in ("soc_bajo", "inicio", "medio", "fin")]
                  + [(cau[id(e)]["previo"] + cau[id(e)]["sin_dato"]) / 60, NOMBRE_CAUSA.get(R.causa_principal(cau[id(e)]), "—")])
        for c_ in range(8, 10): we.cell(i, c_).font = inp
        for c_ in range(12, 19): we.cell(i, c_).font = inp
    ws.append(["Sistema", "Zona", "Marca", "Eventos > 2 min", "T_red (min)", "T_respaldo (min)", "% de respaldo", "Observación", "Sin respaldo: batería sin carga (min)", "Sin respaldo: falla o demora de transición (min)", "% de respaldo sin pérdidas por batería sin carga"])
    for i, r in enumerate(inc, start=2):
        ws.append([r["casa"], zn(r["zona"]), r["marca"], f"=COUNTIFS(Eventos!$B$2:$B${nE},A{i})", f"=SUMIFS(Eventos!$H$2:$H${nE},Eventos!$B$2:$B${nE},A{i})",
                   f"=SUMIFS(Eventos!$J$2:$J${nE},Eventos!$B$2:$B${nE},A{i})", f'=IF(E{i}>0,F{i}/E{i},"Sin eventos")', obs(r),
                   f"=SUMIFS(Eventos!$N$2:$N${nE},Eventos!$B$2:$B${nE},A{i})", f"=SUMIFS(Eventos!$O$2:$O${nE},Eventos!$B$2:$B${nE},A{i})+SUMIFS(Eventos!$P$2:$P${nE},Eventos!$B$2:$B${nE},A{i})+SUMIFS(Eventos!$Q$2:$Q${nE},Eventos!$B$2:$B${nE},A{i})",
                   f'=IF(E{i}>0,(F{i}+I{i})/E{i},"Sin eventos")'])
    mes_txt = f"{MESES[m-1]} {y}"
    wr["A1"] = f"INIC 4 · Disponibilidad operativa (% de respaldo) · {mes_txt.capitalize()}"; wr["A1"].font = Font(name=A, bold=True, size=14, color="1F3A5F")
    wr["A2"] = (f"Meta: ≥ 90 % de disponibilidad operativa promedio de los sistemas. Periodo: {ini.strftime('%d/%m/%Y')} 00:00 a {(fin - dt.timedelta(minutes=1)).strftime('%d/%m/%Y %H:%M')} (hora Colombia). "
                "Alcance: todos los sistemas de Metrum excepto Piloto Promigas, Piloto Huawei, Castellana Real" + "".join(f" y {k} ({v.lower()})" if i == len(excl) - 1 and len(excl) > 0 else f", {k} ({v.lower()})" for i, (k, v) in enumerate(excl.items())) + ".")
    wr["A3"] = f"% Respaldo = (T_red − T_solar) / T_red × 100, consolidado como promedio ponderado por tiempo: suma de T_respaldo ÷ suma de T_red de todos los eventos con T_red mayor a {umbral // 60} minutos."
    for c in ("A2", "A3"): wr[c].font = nf; wr[c].alignment = Alignment(wrap_text=True, vertical="top")
    wr.merge_cells("A2:I2"); wr.merge_cells("A3:I3"); wr.row_dimensions[2].height = 48; wr.row_dimensions[3].height = 32
    def put(r, vals):
        for j, v in enumerate(vals, start=1):
            c = wr.cell(r, j, v); c.font = nf; c.border = bd; c.alignment = Alignment(vertical="center", wrap_text=True)
    blk = [("Sistemas incluidos", f"=COUNTA('Por sistema'!A2:A{nS})", "0"), ("Sistemas con eventos mayores a 2 min", f"=COUNTIF('Por sistema'!D2:D{nS},\">0\")", "0"),
           ("Eventos mayores a 2 min", f"=COUNT(Eventos!H2:H{nE})", "0"), ("T_red total (h)", f"=SUM(Eventos!H2:H{nE})/60", "#,##0.0"), ("T_solar total (h)", f"=SUM(Eventos!I2:I{nE})/60", "#,##0.0"),
           ("T_respaldo total (h)", f"=SUM(Eventos!J2:J{nE})/60", "#,##0.0"), (f"% DE RESPALDO — INIC 4 ({MESES[m-1]})", '=IF(B8>0,B10/B8,"Sin eventos")', "0.0%"), ("¿Cumple la meta ≥ 90 %?", '=IF(ISNUMBER(B11),IF(B11>=0.9,"Sí","No"),"—")', "@")]
    for k, (a_, b_, fm) in enumerate(blk):
        put(5 + k, [a_, b_]); wr.cell(5 + k, 2).number_format = fm
    for c in (1, 2): wr.cell(11, c).fill = PatternFill("solid", fgColor="E3F3E9")
    wr.cell(11, 1).font = bf; wr.cell(11, 2).font = Font(name=A, bold=True, size=13)
    r0 = 15; wr.cell(r0 - 1, 1, "Resumen por zona").font = Font(name=A, bold=True, size=12, color="1F3A5F")
    for j, h_ in enumerate(["Zona", "Sistemas", "Eventos > 2 min", "T_red (h)", "T_respaldo (h)", "% de respaldo", "Sin respaldo: batería sin carga (h)", "Sin respaldo: falla o demora de transición (h)", "% de respaldo sin pérdidas por batería sin carga"], start=1):
        c = wr.cell(r0, j, h_); c.font = hf; c.fill = hfill; c.border = bd; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    zs = []
    for r in inc:
        if zn(r["zona"]) not in zs: zs.append(zn(r["zona"]))
    for k, z in enumerate(zs, start=1):
        r = r0 + k
        put(r, [z, f"=COUNTIF('Por sistema'!$B$2:$B${nS},A{r})", f"=COUNTIF(Eventos!$C$2:$C${nE},A{r})", f"=SUMIFS(Eventos!$H$2:$H${nE},Eventos!$C$2:$C${nE},A{r})/60",
                f"=SUMIFS(Eventos!$J$2:$J${nE},Eventos!$C$2:$C${nE},A{r})/60", f'=IF(D{r}>0,E{r}/D{r},"Sin eventos > 2 min")',
                f"=SUMIFS(Eventos!$N$2:$N${nE},Eventos!$C$2:$C${nE},A{r})/60",
                f"=(SUMIFS(Eventos!$O$2:$O${nE},Eventos!$C$2:$C${nE},A{r})+SUMIFS(Eventos!$P$2:$P${nE},Eventos!$C$2:$C${nE},A{r})+SUMIFS(Eventos!$Q$2:$Q${nE},Eventos!$C$2:$C${nE},A{r}))/60",
                f'=IF(D{r}>0,(E{r}+G{r})/D{r},"Sin eventos > 2 min")'])
        wr.cell(r, 4).number_format = "#,##0.0"; wr.cell(r, 5).number_format = "#,##0.0"; wr.cell(r, 6).number_format = "0.0%"
        wr.cell(r, 7).number_format = "#,##0.00"; wr.cell(r, 8).number_format = "#,##0.00"; wr.cell(r, 9).number_format = "0.0%"
    rt = r0 + len(zs) + 1
    put(rt, ["Total", f"=SUM(B{r0+1}:B{rt-1})", f"=SUM(C{r0+1}:C{rt-1})", f"=SUM(D{r0+1}:D{rt-1})", f"=SUM(E{r0+1}:E{rt-1})", f'=IF(D{rt}>0,E{rt}/D{rt},"Sin eventos")',
              f"=SUM(G{r0+1}:G{rt-1})", f"=SUM(H{r0+1}:H{rt-1})", f'=IF(D{rt}>0,(E{rt}+G{rt})/D{rt},"Sin eventos")'])
    for c in range(1, 10): wr.cell(rt, c).font = bf
    wr.cell(rt, 4).number_format = "#,##0.0"; wr.cell(rt, 5).number_format = "#,##0.0"; wr.cell(rt, 6).number_format = "0.0%"
    wr.cell(rt, 7).number_format = "#,##0.00"; wr.cell(rt, 8).number_format = "#,##0.00"; wr.cell(rt, 9).number_format = "0.0%"
    # ---- composicion del tiempo sin respaldo segun el SOC
    n = rt + 2
    wr.cell(n, 1, "Tiempo sin respaldo según el SOC al apagarse la casa").font = Font(name=A, bold=True, size=12, color="1F3A5F"); n += 1
    for j, h_ in enumerate(["Causa", "Horas", "% del tiempo sin respaldo", "Puntos del T_red", "Sistemas con este tiempo"], start=1):
        c = wr.cell(n, j, h_); c.font = hf; c.fill = hfill; c.border = bd; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wr.row_dimensions[n].height = 30; c0 = n + 1
    filas_c = [("Batería sin carga (SOC ≤ 21 % al apagarse la casa)", "N", ("soc_bajo",)), ("Demora o falla en la transición a OFF-GRID (primeros 3 min, SOC suficiente)", "O", ("inicio",)),
               ("Falla del inversor durante el corte (SOC suficiente)", "P", ("medio",)), ("Demora en el retorno a la red (últimos 90 s, SOC suficiente)", "Q", ("fin",)), ("No atribuible (casa sin energía antes del corte o sin lectura de SOC)", "R", ("previo", "sin_dato"))]
    for k, (lab, col, claves) in enumerate(filas_c):
        r_ = c0 + k
        nsis = len({rr["casa"] for rr, ee in ev if sum(cau[id(ee)][x] for x in claves) > 0})
        put(r_, [lab, f"=SUM(Eventos!{col}2:{col}{nE})/60", f"=IF($B${c0+len(filas_c)}>0,B{r_}/$B${c0+len(filas_c)},0)", f"=IF($B$8>0,B{r_}/$B$8,0)", nsis])
        wr.cell(r_, 5).font = inp; wr.cell(r_, 2).number_format = "#,##0.00"; wr.cell(r_, 3).number_format = "0.0%"; wr.cell(r_, 4).number_format = "0.00%"
    rc = c0 + len(filas_c)
    put(rc, ["Total sin respaldo (T_red − T_respaldo)", f"=SUM(B{c0}:B{rc-1})", f"=SUM(C{c0}:C{rc-1})", f"=SUM(D{c0}:D{rc-1})", ""])
    for c in range(1, 6): wr.cell(rc, c).font = bf
    wr.cell(rc, 2).number_format = "#,##0.00"; wr.cell(rc, 3).number_format = "0.0%"; wr.cell(rc, 4).number_format = "0.00%"
    put(rc + 1, ["Control: T_respaldo + sin respaldo = T_red (h)", f"=B10+B{rc}-B8", "", "", ""]); wr.cell(rc + 1, 2).number_format = "0.000"
    put(rc + 2, ["INIC 4 sin pérdidas por batería sin carga (referencia, no reemplaza al indicador)", f'=IF(B8>0,(B10+B{c0})/B8,"Sin eventos")', "", "", ""])
    wr.cell(rc + 2, 2).number_format = "0.0%"; wr.cell(rc + 2, 1).font = bf; wr.cell(rc + 2, 2).font = bf
    n = rc + 4
    if excl:
        wr.cell(n, 1, "Sistemas fuera del cálculo").font = bf; n += 1
        for k, v in excl.items():
            put(n, [k, v]); wr.merge_cells(start_row=n, start_column=2, end_row=n, end_column=9); n += 1
        n += 1
    wr.cell(n, 1, "Notas").font = bf; n += 1
    parc = [r["casa"] for r in inc if obs(r)]
    notas = [f"Fuente: telemetría de Metrum (eventos de pérdida y restablecimiento de tensión del medidor de red y del medidor solar), consultada el {hoy.strftime('%d/%m/%Y')}.",
             "Evento: cortes de red separados por menos de 15 min se cuentan como un solo evento. T_red = suma del tiempo sin tensión en el medidor de red dentro del evento; T_solar = tiempo sin tensión en el medidor solar entre 2 min antes y 15 min después del evento; T_respaldo = máx(0; T_red − T_solar).",
             ("Los sistemas que ingresaron a operar en el mes o están solo con batería (" + ", ".join(parc) + ") cuentan con los eventos posteriores a su primer dato; ver la columna Observación de la hoja Por sistema.") if parc else "Todos los sistemas incluidos tienen datos completos del mes.",
             "Causa del tiempo sin respaldo (regla del 05/10/2026): para cada apagón de la casa durante un corte de red se toma el SOC del inversor en ese momento (lectura cada 15 min). SOC ≤ 21 % = batería sin carga (consumo del cliente). Con SOC mayor se clasifica por el momento del apagón: primeros 3 min del corte = demora o falla en la transición a OFF-GRID; últimos 90 s = demora en el retorno a la red; entre ambos = falla durante el corte. Casa ya sin energía antes del corte o sin lectura de SOC = no atribuible.",
             "El \"INIC 4 sin pérdidas por batería sin carga\" es solo una referencia para separar la falla del equipo del consumo del cliente; el indicador oficial sigue siendo el % de respaldo definido arriba.",
             "Los valores en azul de la hoja Eventos son datos medidos de Metrum (o derivados de su SOC); el resto son fórmulas. Las zonas sin eventos mayores a 2 minutos se muestran como tales."]
    for t in notas:
        c = wr.cell(n, 1, t); c.font = nf; c.alignment = Alignment(wrap_text=True, vertical="top"); wr.merge_cells(start_row=n, start_column=1, end_row=n, end_column=9); wr.row_dimensions[n].height = max(30, 15 * (len(t) // 170 + 1)); n += 1
    for w in (we, ws):
        for c in w[1]: c.font = hf; c.fill = hfill; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center"); c.border = bd
        w.row_dimensions[1].height = 64; w.freeze_panes = "B2"
    for row in we.iter_rows(min_row=2):
        for c in row:
            if c.font != inp: c.font = nf
            c.border = bd
        row[4].number_format = "dd/mm/yyyy hh:mm:ss"; row[5].number_format = "dd/mm/yyyy hh:mm:ss"
        for k in (7, 8, 9): row[k].number_format = "#,##0.00"
        row[10].number_format = "0.0%"
        for k in range(12, 18): row[k].number_format = "#,##0.00"
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = nf; c.border = bd
        row[4].number_format = "#,##0.00"; row[5].number_format = "#,##0.00"; row[6].number_format = "0.0%"
        row[8].number_format = "#,##0.00"; row[9].number_format = "#,##0.00"; row[10].number_format = "0.0%"
    for w, wd in ((we, [6, 12, 24, 11, 20, 20, 10, 12, 12, 13, 13, 11, 12, 15, 17, 14, 16, 14, 22]), (ws, [12, 24, 11, 14, 13, 15, 14, 48, 17, 19, 19]), (wr, [58, 14, 18, 14, 16, 16, 18, 20, 20])):
        for i, x in enumerate(wd, start=1): w.column_dimensions[L(i)].width = x
    we.auto_filter.ref = f"A1:S{nE}"; ws.auto_filter.ref = f"A1:K{nS}"
    wb.save(args.out)
    # recalcula con LibreOffice si esta disponible (para que los valores se vean sin abrir en Excel)
    so = shutil.which("soffice")
    if so:
        try:
            tmp = tempfile.mkdtemp()
            subprocess.run([so, "--headless", "--convert-to", "xlsx", "--outdir", tmp, args.out], check=True, capture_output=True, timeout=120)
            shutil.move(os.path.join(tmp, os.path.basename(args.out)), args.out)
        except Exception as e:
            print("Aviso: no se pudo recalcular con LibreOffice:", e, file=sys.stderr)
    # resumen en consola (mismo calculo que las formulas del Excel)
    OR = sum(e["or_s"] for r, e in ev); BK = sum(max(0, e["or_s"] - e["cl_s"]) for r, e in ev)
    pct = 100 * BK / OR if OR else None
    print(f"EXCEL: {args.out}")
    print(f"MES: {mes_txt}")
    print(f"SISTEMAS: {len(inc)} incluidos ({len({r['casa'] for r, e in ev})} con eventos > {umbral//60} min); excluidos: {', '.join(excl) or 'ninguno'}")
    print(f"EVENTOS: {len(ev)}; T_red {OR/3600:.1f} h; T_respaldo {BK/3600:.1f} h")
    print(f"INDICADOR: {pct:.1f} %" if pct is not None else "INDICADOR: sin eventos")
    CT = {k: sum(cau[id(e)][k] for r, e in ev) for k in R.CAUSAS}
    print("CAUSAS (h sin respaldo): " + "; ".join(f"{R.CAUSAS[k]}: {v/3600:.2f}" for k, v in CT.items() if v > 0))
    print(f"INDICADOR_SIN_BATERIA_SIN_CARGA (referencia): {100*(BK+CT['soc_bajo'])/OR:.1f} %" if OR else "INDICADOR_SIN_BATERIA_SIN_CARGA: sin eventos")
    print("POR_ZONA:")
    for z in zs:
        o = sum(e["or_s"] for r, e in ev if zn(r["zona"]) == z); k = sum(max(0, e["or_s"] - e["cl_s"]) for r, e in ev if zn(r["zona"]) == z)
        n_s = sum(1 for r in inc if zn(r["zona"]) == z)
        print(f"  {z}: {100*k/o:.1f} % ({n_s} sistemas)" if o else f"  {z}: sin eventos > {umbral//60} min ({n_s} sistemas)")


if __name__ == "__main__":
    main()
